import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

BASE_DIR = r"c:\Users\DELL\Downloads\phase wise example-20261022T103220Z-1-001\IMovies\SkillWallet_Submissions"
os.makedirs(BASE_DIR, exist_ok=True)

epics = {
    "01_Project_Architecture": [
        ("01_Technical_Architecture.docx", "Technical Architecture", 1, "Project Architecture"),
        ("02_ER_Diagram.docx", "ER-Diagram & Data Relationships", 2, "Project Architecture"),
        ("03_Features.docx", "System Features & Specifications", 3, "Project Architecture"),
        ("04_Roles_And_Responsibilities.docx", "Roles and Responsibilities", 4, "Project Architecture"),
        ("05_User_Flow.docx", "User Flow & Booking Journey", 5, "Project Architecture"),
        ("06_MVC_Pattern.docx", "MVC Architectural Pattern", 6, "Project Architecture")
    ],
    "02_Project_Setup_And_Configuration": [
        ("07_Creating_Project_Folder.docx", "Creating Project Folder & Structure", 7, "Project Setup and Configuration"),
        ("08_Client_Setup.docx", "Client Setup & Vite React Tooling", 8, "Project Setup and Configuration"),
        ("09_Server_Setup.docx", "Server Setup & Node/Express Config", 9, "Project Setup and Configuration")
    ],
    "03_Backend_Development": [
        ("10_Backend_Structure.docx", "Backend Directory Structure", 10, "Backend Development"),
        ("11_Backend_Development_And_Explanation.docx", "Backend Development & API Explanation", 11, "Backend Development")
    ],
    "04_Database_Development": [
        ("12_Configure_Mongoose.docx", "Configure Mongoose Connection", 12, "Database Development"),
        ("13_Create_Database_Connection.docx", "Create Database Connection & Resilience", 13, "Database Development"),
        ("14_Create_Schema_And_Models.docx", "Create Schema and Models", 14, "Database Development")
    ],
    "05_Frontend_Development": [
        ("15_Frontend_Structure.docx", "Frontend Structure & Architecture", 15, "Frontend Development"),
        ("16_Frontend_Development_And_Explanation.docx", "Frontend Development & Logic Explanation", 16, "Frontend Development")
    ],
    "06_Project_Execution": [
        ("17_Steps_For_Project_Execution.docx", "Steps for Project Execution", 17, "Project Execution"),
        ("18_Output_Screenshots.docx", "Output Screenshots & Visual Proofs", 18, "Project Execution"),
        ("19_Demo_Video.docx", "Demo Video Walkthrough & Script", 19, "Project Execution")
    ]
}

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def style_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    run = h.runs[0]
    if level == 1:
        run.font.size = Pt(18)
        run.font.color.rgb = RGBColor(229, 9, 20)
        run.font.bold = True
    elif level == 2:
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(30, 30, 30)
        run.font.bold = True
    elif level == 3:
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(70, 70, 70)
        run.font.bold = True
    return h

def add_header_banner(doc, title, epic_name, task_num):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "141414")
    
    p1 = cell.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_sub = p1.add_run(f"iMovies — Movie Ticket Booking Made Simple\nSKILLWALLET SUBMISSION: {epic_name.upper()} | TASK #{task_num}\n")
    r_sub.font.size = Pt(9)
    r_sub.font.color.rgb = RGBColor(229, 9, 20)
    r_sub.font.bold = True
    
    p2 = cell.add_paragraph()
    r_title = p2.add_run(title)
    r_title.font.size = Pt(17)
    r_title.font.color.rgb = RGBColor(255, 255, 255)
    r_title.font.bold = True
    doc.add_paragraph()

def add_callout(doc, text, title="Summary / Key Note:"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F9F0F0")
    
    p = cell.paragraphs[0]
    r_t = p.add_run(f"{title} ")
    r_t.font.bold = True
    r_t.font.color.rgb = RGBColor(229, 9, 20)
    
    r_body = p.add_run(text)
    r_body.font.color.rgb = RGBColor(40, 40, 40)
    doc.add_paragraph()

def add_code_block(doc, code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F4F4F4")
    
    p = cell.paragraphs[0]
    run = p.add_run(code_text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(30, 30, 30)
    doc.add_paragraph()

print("Generator scaffolding ready")
