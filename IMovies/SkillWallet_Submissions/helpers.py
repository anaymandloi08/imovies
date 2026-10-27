import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def style_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    run = h.runs[0]
    if level == 1:
        run.font.size = Pt(20)
        run.font.color.rgb = RGBColor(229, 9, 20) # Cinema Red
        run.font.bold = True
    elif level == 2:
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(30, 30, 30)
        run.font.bold = True
    elif level == 3:
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(80, 80, 80)
        run.font.bold = True
    return h

def add_header_banner(doc, title, epic_name, task_num):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "141414")
    
    p1 = cell.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_sub = p1.add_run(f"iMovies — Movie Ticket Booking Made Simple\nSKILLWALLET SUBMISSION: {epic_name.upper()} | TASK #{task_num}\n")
    r_sub.font.size = Pt(9)
    r_sub.font.color.rgb = RGBColor(229, 9, 20)
    r_sub.font.bold = True
    
    p2 = cell.add_paragraph()
    r_title = p2.add_run(title)
    r_title.font.size = Pt(18)
    r_title.font.color.rgb = RGBColor(255, 255, 255)
    r_title.font.bold = True
    
    doc.add_paragraph() # Spacing

def add_code_block(doc, code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F5F5F5")
    
    p = cell.paragraphs[0]
    run = p.add_run(code_text)
    run.font.name = "Consolas"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(40, 40, 40)
    doc.add_paragraph()

print("Helper script ready")
